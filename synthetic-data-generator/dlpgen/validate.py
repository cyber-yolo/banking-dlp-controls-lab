"""Validators.

validate_output: every generated file exists, matches its manifest hash and carries the markers.
validate_repo:   no real-looking sensitive data or secrets in versioned text files (pre-commit / CI).
"""

from __future__ import annotations

import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from dlpgen import values as v
from dlpgen.generate import sha256_file

TEXT_EXTENSIONS = {".md", ".yaml", ".yml", ".csv", ".json", ".py", ".ps1", ".xml", ".txt", ".toml", ".cfg", ".ini"}
SKIP_DIRS = {".git", ".venv", "venv", "output", "evidence", "lab-exports", "__pycache__", ".pytest_cache", "node_modules"}
IBAN_LENGTHS = {"CH": 21, "LI": 21, "LU": 20, "DE": 22, "FR": 27, "BE": 16, "NL": 18, "AT": 20, "IT": 27, "ES": 24, "GB": 22, "IE": 22, "PT": 25, "MC": 27}
EXAMPLE_EMAIL_SUFFIXES = (".example", ".test", ".invalid", ".localhost", "example.com", "example.org", "example.net")
ALLOWED_TENANT_LABELS = {"contoso", "fabrikam"}

RE_IBAN = re.compile(r"\b[A-Z]{2}\d{2}(?: ?[A-Z0-9]){10,30}")
RE_PAN = re.compile(r"(?<![\w-])(?:\d[ -]?){12,18}\d(?![\w-])")
RE_PRIVATE_KEY = re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")
RE_AWS_KEY = re.compile(r"\b(?:AKIA|ASIA)[0-9A-Z]{16}\b")
RE_EMAIL = re.compile(r"\b[A-Za-z0-9._%+-]+@((?:[A-Za-z0-9-]+\.)+[A-Za-z]{2,})\b")
RE_TENANT = re.compile(r"\b([a-z0-9-]+)\.onmicrosoft\.com\b", re.IGNORECASE)


@dataclass
class Finding:
    path: str
    line: int
    rule: str
    detail: str

    def __str__(self) -> str:
        return f"{self.path}:{self.line}: [{self.rule}] {self.detail}"


def _mask(value: str) -> str:
    return value[:4] + "..." + value[-2:] if len(value) > 8 else "***"


def scan_text(text: str, path: str = "<text>") -> list[Finding]:
    findings = []
    for lineno, line in enumerate(text.splitlines(), start=1):
        for match in RE_IBAN.finditer(line):
            compact = match.group(0).replace(" ", "")
            length = IBAN_LENGTHS.get(compact[:2])
            if not length or len(compact) < length:
                continue
            candidate = compact[:length]
            if v.iban_is_valid(candidate) and v.SYNTH_TAG not in candidate:
                findings.append(Finding(path, lineno, "iban", f"checksum-valid IBAN without {v.SYNTH_TAG} marker: {_mask(candidate)}"))
        for match in RE_PAN.finditer(line):
            digits = re.sub(r"\D", "", match.group(0))
            if 13 <= len(digits) <= 19 and v.luhn_is_valid(digits) and digits not in v.ALL_TEST_PANS:
                findings.append(Finding(path, lineno, "pan", f"Luhn-valid card number not in test list: {_mask(digits)}"))
        if RE_PRIVATE_KEY.search(line):
            findings.append(Finding(path, lineno, "private-key", "private key block"))
        for match in RE_AWS_KEY.finditer(line):
            if match.group(0) != v.AWS_EXAMPLE_ACCESS_KEY:
                findings.append(Finding(path, lineno, "aws-key", f"AWS access key: {_mask(match.group(0))}"))
        for match in RE_EMAIL.finditer(line):
            domain = match.group(1).lower()
            if not domain.endswith(EXAMPLE_EMAIL_SUFFIXES):
                findings.append(Finding(path, lineno, "email", f"non-example email domain: {domain}"))
        for match in RE_TENANT.finditer(line):
            if match.group(1).lower() not in ALLOWED_TENANT_LABELS:
                findings.append(Finding(path, lineno, "tenant", "tenant domain"))
    return findings


def validate_repo(root: Path) -> list[Finding]:
    root = Path(root).resolve()
    findings = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in TEXT_EXTENSIONS:
            continue
        if any(part in SKIP_DIRS for part in path.relative_to(root).parts):
            continue
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            continue
        findings += scan_text(text, str(path.relative_to(root)))
    return findings


def _extract_text(path: Path, fmt: str) -> str:
    if fmt in {"txt", "csv", "json"}:
        return path.read_text(encoding="utf-8")
    if fmt == "docx":
        from docx import Document

        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
        parts += [p.text for s in d.sections for p in s.header.paragraphs + s.footer.paragraphs]
        return "\n".join(parts)
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        return "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    if fmt == "pdf":
        return path.read_bytes().decode("latin-1")
    if fmt == "png":
        from PIL import Image

        with Image.open(path) as img:
            return " ".join(str(val) for val in img.info.values())
    if fmt == "zip":
        with zipfile.ZipFile(path) as zf:
            return "\n".join(zf.read(n).decode("utf-8") for n in zf.namelist())
    raise ValueError(f"unsupported format {fmt}")


def validate_output(manifest_path: Path) -> list[str]:
    manifest_path = Path(manifest_path)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    base = manifest_path.parent
    errors = []
    for entry in manifest["files"]:
        path = base / entry["file_name"]
        if not path.exists():
            errors.append(f"{entry['file_id']}: missing file {entry['file_name']}")
            continue
        if sha256_file(path) != entry["sha256"]:
            errors.append(f"{entry['file_id']}: hash mismatch")
        text = _extract_text(path, entry["format"])
        for marker in v.MARKERS:
            if marker not in text:
                errors.append(f"{entry['file_id']}: marker missing: {marker}")
    return errors
