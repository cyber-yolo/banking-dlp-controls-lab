"""Renderers: one document model, many file formats, deterministic bytes.

Binary office formats are ZIP containers with timestamps. They are normalised
(fixed entry dates, fixed document properties) so that the same seed produces the
same SHA-256 on every run, which keeps regression evidence comparable.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from dlpgen import values as v
from dlpgen.records import Document

ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def _fixed_dt(fixed_iso: str) -> datetime:
    return datetime.fromisoformat(fixed_iso.replace("Z", "+00:00")).astimezone(timezone.utc).replace(tzinfo=None)


def _normalise_zip(path: Path, fixed_iso: str) -> None:
    stamp = _fixed_dt(fixed_iso).strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            if name == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{stamp}\g<2>", text)
                text = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{stamp}\g<2>", text)
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def render_txt(doc: Document, path: Path, fixed_iso: str) -> None:
    path.write_text("\n".join(doc.lines()) + "\n", encoding="utf-8")


def _csv_text(doc: Document) -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow([f"# {v.BANNER}"])
    writer.writerow([doc.title])
    if doc.context:
        writer.writerow([doc.context])
    for section in doc.sections:
        writer.writerow([section.heading])
        writer.writerow(section.headers)
        writer.writerows(section.rows)
    for note in doc.notes:
        writer.writerow([note])
    writer.writerow([f"# {v.BANNER}"])
    return out.getvalue()


def render_csv(doc: Document, path: Path, fixed_iso: str) -> None:
    path.write_text(_csv_text(doc), encoding="utf-8")


def render_json(doc: Document, path: Path, fixed_iso: str) -> None:
    payload = {
        "_marker": v.BANNER,
        "title": doc.title,
        "context": doc.context,
        "sections": [
            {"heading": s.heading, "records": [dict(zip(s.headers, row)) for row in s.rows]} for s in doc.sections
        ],
        "notes": doc.notes,
        "_marker_end": v.BANNER,
    }
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def render_docx(doc: Document, path: Path, fixed_iso: str) -> None:
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.sections[0].header.paragraphs[0].text = v.BANNER
    d.sections[0].footer.paragraphs[0].text = v.BANNER
    d.add_paragraph(v.BANNER)
    d.add_heading(doc.title, level=1)
    if doc.context:
        d.add_paragraph(doc.context)
    for section in doc.sections:
        d.add_heading(section.heading, level=2)
        table = d.add_table(rows=1, cols=len(section.headers))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, section.headers):
            cell.text = header
        for row in section.rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
    for note in doc.notes:
        d.add_paragraph(note)
    d.add_paragraph(v.BANNER)
    props = d.core_properties
    fixed = _fixed_dt(fixed_iso)
    props.title, props.subject, props.keywords = doc.title, v.BANNER, "SYNTHETIC;DLP-TEST"
    props.author = props.last_modified_by = "dlpgen"
    props.created = props.modified = fixed
    props.revision = 1
    d.save(path)
    _normalise_zip(path, fixed_iso)


def render_xlsx(doc: Document, path: Path, fixed_iso: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "SYNTHETIC"
    ws.append([v.BANNER])
    ws.append([doc.title])
    if doc.context:
        ws.append([doc.context])
    for section in doc.sections:
        ws.append([section.heading])
        ws.append(section.headers)
        for row in section.rows:
            ws.append(row)
    for note in doc.notes:
        ws.append([note])
    ws.append([v.BANNER])
    fixed = _fixed_dt(fixed_iso)
    wb.properties.creator = wb.properties.lastModifiedBy = "dlpgen"
    wb.properties.title, wb.properties.subject = doc.title, v.BANNER
    wb.properties.created = wb.properties.modified = fixed
    wb.save(path)
    _normalise_zip(path, fixed_iso)


def _wrap(line: str, width: int = 110) -> list[str]:
    if len(line) <= width:
        return [line]
    return [line[i:i + width] for i in range(0, len(line), width)]


def render_pdf(doc: Document, path: Path, fixed_iso: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1, pageCompression=0)
    c.setTitle(doc.title)
    c.setSubject(v.BANNER)
    c.setAuthor("dlpgen")
    c.setCreator("dlpgen")
    width, height = A4
    lines = [piece for line in doc.lines() for piece in _wrap(line)]

    def page_frame():
        c.setFont("Helvetica-Bold", 8)
        c.drawString(40, height - 25, v.BANNER)
        c.drawString(40, 20, v.BANNER)
        c.setFont("Helvetica", 8)

    page_frame()
    y = height - 50
    for line in lines:
        if y < 40:
            c.showPage()
            page_frame()
            y = height - 50
        c.drawString(40, y, line)
        y -= 11
    c.save()


def render_png(doc: Document, path: Path, fixed_iso: str) -> None:
    from PIL import Image, ImageDraw, ImageFont
    from PIL.PngImagePlugin import PngInfo

    lines = [piece for line in doc.lines() for piece in _wrap(line, 140)]
    img = Image.new("RGB", (1200, 30 + 16 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    for i, line in enumerate(lines):
        draw.text((10, 10 + 16 * i), line, fill="black", font=font)
    info = PngInfo()
    info.add_text("Comment", v.BANNER)
    info.add_text("Title", doc.title)
    img.save(path, pnginfo=info)


def render_zip(doc: Document, path: Path, fixed_iso: str) -> None:
    stem = path.stem
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in ((f"{stem}.txt", "\n".join(doc.lines()) + "\n"), (f"{stem}.csv", _csv_text(doc))):
            info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, text.encode("utf-8"))


RENDERERS = {
    "txt": render_txt,
    "csv": render_csv,
    "json": render_json,
    "docx": render_docx,
    "xlsx": render_xlsx,
    "pdf": render_pdf,
    "png": render_png,
    "zip": render_zip,
}
